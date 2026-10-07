from pathlib import Path

import openpyxl
import pytest

from app.database import SessionLocal
from app.models import Student, User
from app.security import verify_password
from scripts.import_students import import_students

SAMPLES = Path(__file__).resolve().parent.parent / "sample_data"
ORIGINAL = SAMPLES / "student_dataset_10_rows_ORIGINAL.xlsx"
FIXED = SAMPLES / "student_dataset_10_rows_FIXED.xlsx"


def test_original_sheet_is_rejected_because_every_register_number_is_the_same():
    with SessionLocal() as db:
        rep = import_students(db, ORIGINAL)
        assert not rep.ok and rep.rows_read == 10
        assert all("appears in rows" in m for _, m in rep.errors)
        assert db.query(User).count() == 0  # all-or-nothing: nothing written


def test_fixed_sheet_imports_and_sets_ddmmyyyy_passwords():
    with SessionLocal() as db:
        rep = import_students(db, FIXED)
        assert rep.ok and rep.created == 10 and rep.updated == 0
        st = db.query(Student).filter_by(register_no="510423243001").one()
        assert st.user.full_name == "Arun Kumar" and st.department.code == "AIDS"
        assert st.department.name == "Artificial Intelligence & Data Science"
        assert (st.semester, st.batch, st.admission_year, st.section) == (5, "2023-2027", 2023, None)  # "NIL" -> None
        assert st.dob.isoformat() == "2006-05-14" and st.mobile == "9000000001" and st.gender == "Male"
        assert st.user.must_change_password is True
        assert verify_password("14052006", st.user.password_hash)  # date of birth, ddmmyyyy
        assert not verify_password("1405", st.user.password_hash)


def test_sensitive_columns_are_never_stored():
    with SessionLocal() as db:
        rep = import_students(db, FIXED)
        assert {"Aadhar Card Number", "Religion", "Community", "Blood Group", "Father Name"} <= set(rep.sensitive_columns_ignored)
        columns = {c.name for c in Student.__table__.columns}
        assert not any(w in " ".join(columns) for w in ("aadhar", "aadhaar", "religion", "community", "blood", "father", "mother"))
        dumped = " ".join(str(v) for row in db.execute(Student.__table__.select()) for v in row)
        for secret in ("000000000001", "Ramesh Kumar", "O+", "Hindu", "9000000011"):  # Aadhaar, father, blood group...
            assert secret not in dumped


def test_dry_run_writes_nothing():
    with SessionLocal() as db:
        rep = import_students(db, FIXED, dry_run=True)
        assert rep.ok and rep.rows_read == 10 and db.query(User).count() == 0


def test_reimport_updates_details_but_keeps_passwords_sections_and_counsellor():
    with SessionLocal() as db:
        import_students(db, FIXED)
        st = db.query(Student).filter_by(register_no="510423243001").one()
        st.section, st.user.password_hash = "A", "$2b$04$kept"
        db.commit()
        rep = import_students(db, FIXED)
        assert rep.created == 0 and rep.updated == 10
        st = db.query(Student).filter_by(register_no="510423243001").one()
        assert st.section == "A"                      # not overwritten by "NIL"
        assert st.user.password_hash == "$2b$04$kept"  # password never reset by a re-import


def edit_sheet(tmp_path, mutate):
    wb = openpyxl.load_workbook(FIXED)
    mutate(wb.active)
    path = tmp_path / "edited.xlsx"
    wb.save(path)
    return path


@pytest.mark.parametrize("col,value,fragment", [
    ("Register Number", 123, "12 digits starting with 5104"),
    ("Register Number", 410423243001, "12 digits starting with 5104"),
    ("Date of Birth", "not a date", "not a date"),
    ("Semester", 9, "semester"),
    ("Batch", "2023", "2023-2027"),
    ("Name", None, "'name' is empty"),
    ("Gender", "robot", "gender"),
])
def test_bad_rows_are_reported_with_row_numbers_and_nothing_is_written(tmp_path, col, value, fragment):
    def mutate(ws):
        header = [c.value for c in ws[1]]
        ws.cell(row=4, column=header.index(col) + 1).value = value
    path = edit_sheet(tmp_path, mutate)
    with SessionLocal() as db:
        rep = import_students(db, path)
        assert not rep.ok and any(r == 4 and fragment in m for r, m in rep.errors), rep.errors
        assert db.query(User).count() == 0


def test_missing_required_column(tmp_path):
    def mutate(ws):
        header = [c.value for c in ws[1]]
        ws.delete_cols(header.index("Date of Birth") + 1)
    with SessionLocal() as db:
        rep = import_students(db, edit_sheet(tmp_path, mutate))
        assert not rep.ok and "missing columns" in rep.errors[0][1] and "dob" in rep.errors[0][1]


def test_mismatched_register_number_is_only_a_warning(tmp_path):
    def mutate(ws):
        header = [c.value for c in ws[1]]
        ws.cell(row=2, column=header.index("Batch") + 1).value = "2024-2028"
    with SessionLocal() as db:
        rep = import_students(db, edit_sheet(tmp_path, mutate))
        assert rep.ok and any("batch start year" in m for _, m in rep.warnings)


def test_assign_to_counsellor_gives_every_unassigned_student_a_counsellor():
    from app.enums import Role
    from app.models import Department, User
    with SessionLocal() as db:
        d = Department(code="AIDS", name="x")
        db.add(d)
        db.flush()
        db.add(User(username="ANITHA-STAFF", full_name="Anitha", role=Role.COUNSELLOR, department_id=d.id,
                    password_hash="x", must_change_password=True))
        db.commit()
        assert not import_students(db, FIXED, assign_to="nobody").ok
        rep = import_students(db, FIXED, assign_to="anitha-staff")
        assert rep.ok and rep.assigned == 10
        assert all(s.counsellor.username == "ANITHA-STAFF" for s in db.query(Student).all())
